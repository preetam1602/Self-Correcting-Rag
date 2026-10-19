"""Parse PDF, HTML, DOCX and Markdown into ordered text blocks with section and page."""
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class Block:
    text: str
    section: str | None     # heading trail, e.g. "Leave policy > Notice periods"
    page: int | None        # PDF only


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    blocks: list[Block]


def _section(trail: list[str]) -> str | None:
    return " > ".join(trail) or None


def _push_heading(trail: list[str], level: int, heading: str) -> list[str]:
    return trail[:level - 1] + [heading]


# ---------------------------------------------------------------- Markdown

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def parse_markdown(path: Path) -> ParsedDocument:
    title, trail, blocks, para = None, [], [], []
    in_fence = False

    def flush():
        text = "\n".join(para).strip()
        if text:
            blocks.append(Block(text, _section(trail), None))
        para.clear()

    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if _MD_FENCE.match(line):
            in_fence = not in_fence
        elif not in_fence:
            heading = _MD_HEADING.match(line)
            if heading:
                flush()
                level, text = len(heading.group(1)), heading.group(2)
                trail = _push_heading(trail, level, text)
                if title is None and level == 1:
                    title = text
                continue
            if not line.strip():
                flush()
                continue
        para.append(line)
    flush()
    return ParsedDocument(title or path.stem, blocks)


# ---------------------------------------------------------------- HTML

_HTML_HEADINGS = ["h1", "h2", "h3", "h4", "h5", "h6"]
_HTML_BLOCKS = ["p", "li", "pre", "blockquote", "tr", "dt", "dd", "figcaption", "caption",
                "div", "section", "article"]
_HTML_DROP = ["script", "style", "noscript", "template", "nav"]


def parse_html(path: Path) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_bytes(), "lxml")
    for tag in soup(_HTML_DROP):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else ""
    root = soup.body or soup
    trail, blocks, first_h1 = [], [], None

    for el in root.find_all(_HTML_HEADINGS + _HTML_BLOCKS):
        if el.name in _HTML_HEADINGS:
            text = el.get_text(" ", strip=True)
            if text:
                trail = _push_heading(trail, int(el.name[1]), text)
                if first_h1 is None and el.name == "h1":
                    first_h1 = text
            continue
        if el.find(_HTML_HEADINGS + _HTML_BLOCKS):     # only leaf blocks carry text
            continue
        if el.name == "tr":
            text = " | ".join(c.get_text(" ", strip=True) for c in el.find_all(["th", "td"]))
        else:
            text = el.get_text(" ", strip=True)
        if text.strip(" |"):
            blocks.append(Block(text, _section(trail), None))
    return ParsedDocument(title or first_h1 or path.stem, blocks)


# ---------------------------------------------------------------- DOCX

_DOCX_HEADING = re.compile(r"Heading (\d)")


def parse_docx(path: Path) -> ParsedDocument:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(str(path))
    title = (doc.core_properties.title or "").strip() or None
    trail, blocks = [], []

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style is not None else ""
            if style == "Title":
                title = title or text
                continue
            heading = _DOCX_HEADING.match(style)
            if heading:
                trail = _push_heading(trail, int(heading.group(1)), text)
                continue
            blocks.append(Block(text, _section(trail), None))
        elif child.tag == qn("w:tbl"):
            for row in Table(child, doc).rows:
                cells = []
                for cell in row.cells:                  # merged cells repeat; keep one
                    text = cell.text.strip()
                    if text and (not cells or cells[-1] != text):
                        cells.append(text)
                if cells:
                    blocks.append(Block(" | ".join(cells), _section(trail), None))
    return ParsedDocument(title or path.stem, blocks)


# ---------------------------------------------------------------- PDF

_NUMERIC_LINE = re.compile(r"^[\d\s.,%()+\-\u2212]*$")       # page numbers, chart axis ticks
_SMALL_CAPS = re.compile(r"\b([A-Z]) ([A-Z]{2,})\b")           # "C ONCLUSION", "R AG"
_SPACED_HYPHEN = re.compile(r"(\w) -\n?(\w)")                   # "SELF -RAG", "SELF -\nRAG"
_LINE_BREAK_HYPHEN = re.compile(r"(\w+)-\n(\w+)")               # "Gen-\neration"
_WORD = re.compile(r"\w+(?:-\w+)*")


def _repeated_lines(pages: list[str]) -> set[str]:
    """Running headers and footers: short lines found on at least half the pages (3+ pages)."""
    if len(pages) < 3:
        return set()
    counts: dict[str, int] = {}
    for text in pages:
        for line in {line.strip() for line in text.splitlines() if 0 < len(line.strip()) <= 80}:
            counts[line] = counts.get(line, 0) + 1
    return {line for line, n in counts.items() if n >= len(pages) / 2}


def pdf_vocabulary(pages: list[str]) -> set[str]:
    """Lower-cased words of the document, leaving out small-caps fragments like "ONCLUSION"."""
    return {w.lower() for text in pages
            for w in _WORD.findall(_SMALL_CAPS.sub(" ", unicodedata.normalize("NFKC", text)))}


def clean_pdf_text(text: str, drop: set[str], vocabulary: set[str]) -> str:
    """Undo common PDF extraction artefacts, using the document's own words as a dictionary."""
    text = unicodedata.normalize("NFKC", text)              # ligatures: "ﬂ" -> "fl"
    lines = [line for line in text.splitlines()
             if line.strip() not in drop and not _NUMERIC_LINE.match(line.strip())]
    text = "\n".join(lines)

    def small_caps(m: re.Match) -> str:
        # Join when the result is a word in this document, or the tail on its own is not
        # ("C ONCLUSION"), but keep labels before real words ("B EXPERIMENTAL", "A NEW").
        joined, tail = m.group(1) + m.group(2), m.group(2).lower()
        return joined if joined.lower() in vocabulary or tail not in vocabulary else m.group(0)

    def line_break_hyphen(m: re.Match) -> str:
        left, right = m.group(1), m.group(2)
        hyphenated = f"{left}-{right}".lower()
        if right[0].isupper() or (hyphenated in vocabulary
                                  and (left + right).lower() not in vocabulary):
            return f"{left}-{right}"          # a real compound, e.g. "Retrieval-Augmented"
        return left + right                   # a word split across lines, e.g. "Gen-eration"

    text = _SMALL_CAPS.sub(small_caps, text)
    text = _SPACED_HYPHEN.sub(r"\1-\2", text)
    return _LINE_BREAK_HYPHEN.sub(line_break_hyphen, text)


def parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    title = ((reader.metadata.title if reader.metadata else None) or "").strip()
    pages = [page.extract_text() or "" for page in reader.pages]
    drop = _repeated_lines(pages)
    vocabulary = pdf_vocabulary(pages)
    blocks = []
    for number, text in enumerate(pages, start=1):
        for para in re.split(r"\n\s*\n", clean_pdf_text(text, drop, vocabulary)):
            if para.strip():
                blocks.append(Block(para.strip(), None, number))
    return ParsedDocument(title or path.stem, blocks)


# ---------------------------------------------------------------- dispatch

PARSERS = {
    ".pdf": parse_pdf,
    ".html": parse_html, ".htm": parse_html,
    ".docx": parse_docx,
    ".md": parse_markdown, ".markdown": parse_markdown,
}


def parse_file(path: Path) -> ParsedDocument:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"unsupported file type: {path.suffix} (supported: {', '.join(PARSERS)})")
    return parser(path)
