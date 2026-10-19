"""Builders for small test documents."""
from pathlib import Path


def write_pdf(path: Path, pages: list[str]) -> Path:
    """Write a minimal valid PDF with one line of Helvetica text per page."""
    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objects = ["<< /Type /Catalog /Pages 2 0 R >>",
               f"<< /Type /Pages /Kids [{kids}] /Count {n} >>",
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out, offsets = b"%PDF-1.4\n", []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    path.write_bytes(out)
    return path


def write_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Contractor Handbook", style="Title")
    doc.add_heading("Notice", level=1)
    doc.add_paragraph("Contractors must give 30 days notice.")
    doc.add_heading("Payment", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Type", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "Invoice", "15"
    doc.save(str(path))
    return path
